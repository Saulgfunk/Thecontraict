"""Generated documents: AI-drafted notices, Word downloads and Excel exports."""

import io
import json
import logging
import re
import uuid
from collections.abc import Iterator
from datetime import date, timedelta
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select

from app import audit, chat, drafting
from app.deps import OrgContext, get_org_context
from app.models import (
    Contract,
    Deadline,
    DeadlineStatus,
    User,
    Workspace,
    WorkspaceMembership,
    WorkspaceRole,
)
from app.routers.contracts import _get_contract

router = APIRouter(prefix="/organizations/{org_id}", tags=["documents"])
log = logging.getLogger(__name__)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sse(payload: dict[str, Any]) -> str:
    return f"data: {json.dumps(payload, default=str)}\n\n"


def _filename(name: str, ext: str) -> str:
    safe = re.sub(r"[^\w\- ]", "", name).strip().replace(" ", "_")[:80] or "document"
    return f"{safe}.{ext}"


# ---------------------------------------------------------------------------
# Notice drafting
# ---------------------------------------------------------------------------


class DraftNoticeIn(BaseModel):
    kind: drafting.NoticeKind
    deadline_id: uuid.UUID | None = None
    sender: str | None = Field(default=None, max_length=500, description="Our side, as signatory")
    instructions: str | None = Field(default=None, max_length=2000)


@router.post(
    "/contracts/{contract_id}/draft-notice",
    response_class=StreamingResponse,
    responses={200: {"content": {"text/event-stream": {}}}},
)
def draft_notice(
    contract_id: uuid.UUID, body: DraftNoticeIn, ctx: OrgContext = Depends(get_org_context)
) -> StreamingResponse:
    """Stream a draft notice letter as Server-Sent Events: ``delta`` events with text,
    then ``done`` with the full text and the clauses it relied on, or ``error``."""
    contract = _get_contract(ctx, contract_id, WorkspaceRole.EDITOR)
    deadline = None
    if body.deadline_id:
        deadline = ctx.db.get(Deadline, body.deadline_id)
        if deadline is None or deadline.contract_id != contract.id:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Deadline not found")
    sources = chat.build_sources(ctx.db, [contract], "", single=True)
    if not sources:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "The contract's document has not been analysed yet"
        )
    messages = drafting.build_request(
        contract, sources[0], body.kind, deadline, body.sender, body.instructions, date.today()
    )
    titles = {contract.id: contract.title}
    audit.record(
        ctx.db,
        organization_id=ctx.org_id,
        workspace_id=contract.workspace_id,
        actor_user_id=ctx.user.id,
        action="contract.notice_drafted",
        entity_type="contract",
        entity_id=contract.id,
        data={
            "kind": body.kind,
            "deadline_id": str(body.deadline_id) if body.deadline_id else None,
        },
    )
    ctx.db.commit()

    def events() -> Iterator[str]:
        try:
            model = chat.get_chat_model()
            for item in model.stream(drafting.DRAFTING_SYSTEM_PROMPT, messages):
                if isinstance(item, str):
                    yield _sse({"type": "delta", "text": item})
                    continue
                text = "".join(b["text"] for b in item.blocks)
                citations = [
                    c
                    for b in item.blocks
                    for c in chat.resolve_citations(b.get("citations", []), sources, titles)
                ]
                yield _sse({"type": "done", "text": text.strip(), "citations": citations})
                return
            raise chat.ChatError("The draft ended unexpectedly")
        except Exception as exc:
            log.exception("Drafting failed")
            detail = str(exc) if isinstance(exc, chat.ChatError) else "The AI service failed"
            yield _sse({"type": "error", "detail": detail})

    return StreamingResponse(
        events(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


class DocxIn(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    text: str = Field(min_length=1, max_length=100_000)


@router.post("/render-docx", response_class=Response)
def render_docx(body: DocxIn, ctx: OrgContext = Depends(get_org_context)) -> Response:
    """Turn (edited) letter text into a Word document."""
    return Response(
        content=drafting.render_docx(body.title, body.text),
        media_type=DOCX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{_filename(body.title, "docx")}"'},
    )


# ---------------------------------------------------------------------------
# Excel exports
# ---------------------------------------------------------------------------


def _accessible_workspaces(ctx: OrgContext, workspace_id: uuid.UUID | None) -> list[Workspace]:
    if workspace_id:
        return [ctx.get_workspace(workspace_id)]
    query = select(Workspace).where(Workspace.organization_id == ctx.org_id)
    if not ctx.is_org_admin:
        query = query.where(
            Workspace.id.in_(
                select(WorkspaceMembership.workspace_id).where(
                    WorkspaceMembership.user_id == ctx.user.id
                )
            )
        )
    return list(ctx.db.scalars(query.order_by(Workspace.name)))


def _xlsx(sheet: str, headers: list[str], rows: list[list[Any]], widths: list[int]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E40AF")
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_response(data: bytes, name: str) -> Response:
    return Response(
        content=data,
        media_type=XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@router.get("/exports/deadlines.xlsx", response_class=Response)
def export_deadlines(
    ctx: OrgContext = Depends(get_org_context),
    workspace_id: uuid.UUID | None = None,
    start: date | None = Query(default=None, alias="from"),
    end: date | None = Query(default=None, alias="to"),
    include_closed: bool = True,
) -> Response:
    workspaces = {w.id: w for w in _accessible_workspaces(ctx, workspace_id)}
    start = start or date.today() - timedelta(days=90)
    query = (
        select(Deadline, Contract)
        .join(Contract, Contract.id == Deadline.contract_id)
        .where(Deadline.workspace_id.in_(workspaces), Deadline.due_date >= start)
        .order_by(Deadline.due_date)
    )
    if end:
        query = query.where(Deadline.due_date <= end)
    if not include_closed:
        query = query.where(Deadline.status == DeadlineStatus.OPEN)
    results = ctx.db.execute(query).all()
    owner_ids = {c.owner_id for _, c in results if c.owner_id}
    users = (
        {u.id: u.email for u in ctx.db.scalars(select(User).where(User.id.in_(owner_ids)))}
        if owner_ids
        else {}
    )
    rows = [
        [
            d.due_date,
            (d.due_date - date.today()).days,
            d.label,
            d.kind.replace("_", " "),
            c.title,
            c.counterparty_name,
            workspaces[d.workspace_id].name,
            "yes" if d.confirmed else "no",
            d.status,
            (d.decision or "").replace("_", " "),
            d.decision_note,
            users.get(c.owner_id) if c.owner_id else None,
            " / ".join(d.derivation),
        ]
        for d, c in results
    ]
    data = _xlsx(
        "Deadlines",
        [
            "Due date",
            "Days left",
            "Deadline",
            "Type",
            "Contract",
            "Counterparty",
            "Workspace",
            "Confirmed",
            "Status",
            "Decision",
            "Decision note",
            "Owner",
            "How calculated",
        ],
        rows,
        [12, 10, 32, 12, 36, 28, 22, 10, 10, 14, 30, 28, 70],
    )
    audit.record(
        ctx.db,
        organization_id=ctx.org_id,
        actor_user_id=ctx.user.id,
        action="export.deadlines",
        entity_type="organization",
        entity_id=ctx.org_id,
        data={"rows": len(rows)},
    )
    ctx.db.commit()
    return _xlsx_response(data, f"deadlines-{date.today().isoformat()}.xlsx")


@router.get("/exports/contracts.xlsx", response_class=Response)
def export_contracts(
    ctx: OrgContext = Depends(get_org_context), workspace_id: uuid.UUID | None = None
) -> Response:
    workspaces = {w.id: w for w in _accessible_workspaces(ctx, workspace_id)}
    contracts = ctx.db.scalars(
        select(Contract)
        .where(Contract.workspace_id.in_(workspaces))
        .order_by(Contract.workspace_id, Contract.title)
    ).all()
    today = date.today()
    rows = []
    for c in contracts:
        upcoming = [
            d for d in c.deadlines if d.status == DeadlineStatus.OPEN and d.due_date >= today
        ]
        nxt = min(upcoming, key=lambda d: d.due_date, default=None)
        term = f"{c.initial_term_amount} {c.initial_term_unit}" if c.initial_term_amount else None
        renewal = (
            f"{c.renewal_term_amount} {c.renewal_term_unit}" if c.renewal_term_amount else None
        )
        rows.append(
            [
                c.title,
                workspaces[c.workspace_id].name,
                c.counterparty_name,
                c.contract_type,
                c.status.replace("_", " "),
                c.effective_date,
                term,
                {True: "yes", False: "no", None: ""}[c.auto_renews],
                renewal,
                c.end_date,
                c.governing_law,
                c.currency,
                float(c.contract_value) if c.contract_value is not None else None,
                nxt.label if nxt else None,
                nxt.due_date if nxt else None,
                "yes" if c.reviewed_at else "no",
            ]
        )
    data = _xlsx(
        "Contracts",
        [
            "Contract",
            "Workspace",
            "Counterparty",
            "Type",
            "Status",
            "Effective date",
            "Initial term",
            "Auto-renews",
            "Renewal term",
            "End date",
            "Governing law",
            "Currency",
            "Value",
            "Next deadline",
            "Next deadline date",
            "Terms reviewed",
        ],
        rows,
        [36, 22, 28, 16, 12, 13, 13, 11, 13, 12, 22, 9, 14, 30, 16, 13],
    )
    audit.record(
        ctx.db,
        organization_id=ctx.org_id,
        actor_user_id=ctx.user.id,
        action="export.contracts",
        entity_type="organization",
        entity_id=ctx.org_id,
        data={"rows": len(rows)},
    )
    ctx.db.commit()
    return _xlsx_response(data, f"contracts-{today.isoformat()}.xlsx")
