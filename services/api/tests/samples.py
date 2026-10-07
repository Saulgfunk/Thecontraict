"""Sample contracts and a deterministic stand-in for the Claude extractor."""

import io
import re
from datetime import date

from app.pipeline import extract as ext

MSA_LINES = [
    "MASTER SERVICES AGREEMENT",
    'This Agreement is made between Acme Supplies Ltd ("Supplier") and Client A Ltd ("Customer").',
    "1. DEFINITIONS",
    "In this Agreement, Services means the services described in Schedule 1.",
    "2. TERM",
    "2.1 This Agreement commences on 1 January 2025 and continues for an initial term of "
    "24 months.",
    "2.2 Thereafter it renews automatically for successive periods of 12 months unless either "
    "party gives the other at least 90 days' written notice of non-renewal before the end of "
    "the then-current term.",
    "3. FEES",
    "3.1 The Customer shall pay a monthly fee of EUR 5,000, invoiced on the first day of each "
    "month starting 1 January 2025 and payable within 30 days.",
    "4. NOTICES",
    "Notices must be sent by registered post to the addresses above and are deemed received "
    "2 business days after posting.",
    "5. GOVERNING LAW",
    "This Agreement is governed by the laws of England and Wales.",
]


def msa_docx(extra: str = "") -> bytes:
    import docx

    document = docx.Document()
    for line in MSA_LINES:
        document.add_paragraph(line)
    if extra:
        document.add_paragraph(extra)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def msa_pdf() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for i, line in enumerate(MSA_LINES):
        if i in (0, 7):  # two pages
            pdf.add_page()
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def blank_pdf(pages: int = 2) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    for _ in range(pages):
        pdf.add_page()
    return bytes(pdf.output())


def _src(refs: list[str], quote: str, confidence: float = 0.95) -> ext.Source:
    return ext.Source(clause_refs=refs, quote=quote, confidence=confidence)


class FakeExtractor:
    """Returns the extraction a good model would produce for MSA_LINES, citing the refs the
    pipeline actually assigned."""

    model = "fake-model"

    def __init__(self) -> None:
        self.calls = 0
        self.ocr_calls = 0
        self.last_xml = ""
        self.last_context = ""

    def extract(self, contract_xml: str, context: str) -> ext.ExtractionOutput:
        self.calls += 1
        self.last_xml, self.last_context = contract_xml, context
        refs = dict(
            (number, ref)
            for ref, number in re.findall(r'<clause ref="(C\d+)" number="([^"]+)"', contract_xml)
        )
        term, renew, fees = refs.get("2.1", "C?"), refs.get("2.2", "C?"), refs.get("3.1", "C?")
        notices, law = refs.get("4", "C?"), refs.get("5", "C?")
        data = ext.ContractExtraction(
            title="Master Services Agreement",
            contract_type="Services",
            summary="Acme supplies services to Client A for 24 months, renewing yearly.",
            parties=[
                ext.Party(name="Acme Supplies Ltd", role="Supplier"),
                ext.Party(name="Client A Ltd", role="Customer"),
            ],
            counterparty_name="Acme Supplies Ltd",
            effective_date=ext.SourcedDate(
                value=date(2025, 1, 1), source=_src([term], "commences on 1 January 2025")
            ),
            end_date=None,
            initial_term=ext.SourcedPeriod(
                value=ext.Period(amount=24, unit="months"),
                source=_src([term], "initial term of 24 months"),
            ),
            auto_renews=ext.SourcedBool(value=True, source=_src([renew], "renews automatically")),
            renewal_term=ext.SourcedPeriod(
                value=ext.Period(amount=12, unit="months"),
                source=_src([renew], "successive periods of 12 months"),
            ),
            governing_law=ext.SourcedText(
                value="England and Wales", source=_src([law], "laws of England and Wales")
            ),
            currency="eur",
            contract_value=None,
            notice_details=ext.SourcedText(
                value="Registered post to the parties' addresses",
                source=_src([notices], "sent by registered post"),
            ),
            date_rules=[
                ext.ExtractedDateRule(
                    rule_type="non_renewal_notice",
                    label="Notice of non-renewal",
                    anchor="term_end",
                    fixed_date=None,
                    offset=ext.Period(amount=90, unit="days"),
                    offset_basis="calendar",
                    direction="before",
                    deemed_receipt_business_days=2,
                    source=_src([renew, notices, "C999"], "at least 90 days' written notice"),
                )
            ],
            payment_terms=[
                ext.ExtractedPaymentTerm(
                    description="Monthly service fee",
                    direction="payable",
                    amount=5000,
                    currency="EUR",
                    frequency="monthly",
                    first_due_date=date(2025, 1, 1),
                    payment_days=30,
                    escalation=None,
                    source=_src([fees], "monthly fee of EUR 5,000"),
                )
            ],
        )
        return ext.ExtractionOutput(
            extraction=data, model=self.model, input_tokens=1000, output_tokens=500
        )

    def transcribe_pdf(self, pdf: bytes) -> list[ext.OcrPage]:
        self.ocr_calls += 1
        return [
            ext.OcrPage(page_number=1, text="\n".join(MSA_LINES[:7])),
            ext.OcrPage(page_number=2, text="\n".join(MSA_LINES[7:])),
        ]
